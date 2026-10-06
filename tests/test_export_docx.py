# -*- coding: utf-8 -*-
"""Word 导出模块的单元测试。

重点验证两件事：
1. 导出顺序 / 乱序（含随机种子可复现）的行为；
2. 藏文 run 真的写入了复杂文本字体（w:rFonts@cs / @eastAsia）和 w:szCs——
   少了这两样，Word 里藏文会显示成方框。
"""

from __future__ import annotations

import io
import re
import zipfile
from datetime import date

import pytest
from docx import Document
from docx.oxml.ns import qn

import models
from export_docx import (
    DEFAULT_TIBETAN_FONT,
    DEFAULT_TIBETAN_SIZE,
    LAYOUT_COLUMNS,
    LAYOUT_TABLE,
    ORDER_RANDOM,
    ORDER_SEQUENTIAL,
    TIBETAN_SIZE_CHOICES,
    ExportOptions,
    _safe_text,
    apply_order,
    build_document,
    build_filename,
    export_to_bytes,
)


# ---------------------------------------------------------------------------
# 工具
# ---------------------------------------------------------------------------
def make_rows(count: int, meaning_prefix: str = "释义") -> list[models.Word]:
    return [
        models.Word(id=i + 1, tibetan=f"ཀ{i}", meaning=f"{meaning_prefix}{i}")
        for i in range(count)
    ]


def body_runs(document: Document) -> list:
    """文档正文里所有 run：先段落，再表格单元格。

    表格排版只用到表格，分栏排版只用到段落，所以两种排版都得看。
    """
    runs = []
    for paragraph in document.paragraphs:
        runs.extend(paragraph.runs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    runs.extend(paragraph.runs)
    return runs


def tibetan_runs(document: Document) -> list:
    """只取正文里含藏文字符的 run（释义 run 用的是中文字体，不参与断言）。"""
    return [run for run in body_runs(document)
            if any(0x0F00 <= ord(ch) <= 0x0FFF for ch in run.text)]


def table_options(**overrides) -> ExportOptions:
    """表格排版的导出设置。

    默认排版已经改成「分栏」，所以专门验证表格行为的用例必须显式指定
    layout=table，否则拿不到 document.tables。
    """
    overrides.setdefault("layout", LAYOUT_TABLE)
    overrides.setdefault("landscape", False)
    return ExportOptions(**overrides)


def document_xml(data: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        return archive.read("word/document.xml").decode("utf-8")


# ---------------------------------------------------------------------------
# 导出顺序
# ---------------------------------------------------------------------------
def test_sequential_keeps_original_order():
    rows = make_rows(10)
    options = table_options(order=ORDER_SEQUENTIAL)
    assert [w.id for w in apply_order(rows, options)] == [w.id for w in rows]


def test_random_shuffle_reorders_rows():
    rows = make_rows(30)
    shuffled = [w.id for w in apply_order(rows, table_options(order=ORDER_RANDOM))]
    assert sorted(shuffled) == [w.id for w in rows]      # 不丢不重
    assert shuffled != [w.id for w in rows]              # 确实打乱了


def test_random_seed_is_reproducible():
    rows = make_rows(30)
    first = [w.id for w in apply_order(rows, table_options(order=ORDER_RANDOM, seed="42"))]
    second = [w.id for w in apply_order(rows, table_options(order=ORDER_RANDOM, seed="42"))]
    third = [w.id for w in apply_order(rows, table_options(order=ORDER_RANDOM, seed="43"))]

    assert first == second       # 同一个种子 -> 完全相同的顺序
    assert first != third        # 不同种子 -> 不同顺序


def test_random_seed_accepts_text():
    rows = make_rows(30)
    first = [w.id for w in apply_order(rows, table_options(order=ORDER_RANDOM, seed="西藏"))]
    second = [w.id for w in apply_order(rows, table_options(order=ORDER_RANDOM, seed="西藏"))]
    assert first == second


def test_empty_seed_gives_different_orders():
    rows = make_rows(60)
    orders = {
        tuple(w.id for w in apply_order(rows, table_options(order=ORDER_RANDOM)))
        for _ in range(5)
    }
    assert len(orders) > 1


def test_sequential_ignores_seed():
    rows = make_rows(10)
    options = table_options(order=ORDER_SEQUENTIAL, seed="42")
    assert [w.id for w in apply_order(rows, options)] == [w.id for w in rows]


# ---------------------------------------------------------------------------
# 文件名
# ---------------------------------------------------------------------------
def test_filename_contains_order_label_and_date():
    day = date(2026, 3, 8)

    sequential = build_filename(table_options(order=ORDER_SEQUENTIAL), today=day)
    assert "顺序" in sequential
    assert "20260308" in sequential
    assert sequential.endswith(".docx")

    random_name = build_filename(table_options(order=ORDER_RANDOM), today=day)
    assert "乱序" in random_name
    assert "20260308" in random_name


def test_filename_uses_title_and_strips_illegal_characters():
    name = build_filename(
        table_options(title='藏语/词汇:测试*表?'), today=date(2026, 1, 1)
    )
    assert "/" not in name and ":" not in name and "*" not in name and "?" not in name
    assert "藏语" in name
    assert "顺序" in name


def test_filename_falls_back_when_title_is_blank():
    name = build_filename(table_options(title="   "), today=date(2026, 1, 1))
    assert name.startswith("藏语词汇")


# ---------------------------------------------------------------------------
# 藏文 run 的字体设置（最容易出问题的地方）
# ---------------------------------------------------------------------------
def test_tibetan_run_sets_complex_script_font():
    document = build_document(
        make_rows(3), table_options(tibetan_font="Jomolhari")
    )
    runs = tibetan_runs(document)
    assert runs, "表格里应当有藏文 run"

    for run in runs:
        rPr = run._element.rPr
        assert rPr is not None
        rFonts = rPr.find(qn("w:rFonts"))
        assert rFonts is not None
        # 复杂文本字体 —— Word 渲染藏文时取的就是这个
        assert rFonts.get(qn("w:cs")) == "Jomolhari"
        assert rFonts.get(qn("w:eastAsia")) == "Jomolhari"
        assert rFonts.get(qn("w:ascii")) == "Jomolhari"
        assert rFonts.get(qn("w:hAnsi")) == "Jomolhari"


def test_tibetan_run_sets_szcs():
    document = build_document(make_rows(2), table_options(tibetan_size_pt=16.0))
    for run in tibetan_runs(document):
        szCs = run._element.rPr.find(qn("w:szCs"))
        assert szCs is not None, "缺少 w:szCs，Word 会用默认字号渲染藏文"
        # w:szCs 的单位是半磅
        assert szCs.get(qn("w:val")) == "32"


def test_tibetan_runs_are_bold():
    """藏文一律加粗——不加粗的话打印出来发虚。

    藏文字体大多没有真正的粗体字重，Word 会合成加粗，视觉上是「更实」。
    注意复杂文本靠的是 w:bCs，只设 w:b 对藏文没用。
    """
    for layout in (LAYOUT_TABLE, LAYOUT_COLUMNS):
        document = build_document(make_rows(2), ExportOptions(layout=layout))
        runs = tibetan_runs(document)
        assert runs, layout
        for run in runs:
            rPr = run._element.rPr
            assert rPr.find(qn("w:b")) is not None, layout
            assert rPr.find(qn("w:bCs")) is not None, layout


def test_tibetan_default_size_is_large_enough_to_print():
    """默认字号要比中文正文大不少，否则打印出来看不清。"""
    options = ExportOptions()
    assert options.tibetan_size_pt == DEFAULT_TIBETAN_SIZE == 20.0
    assert options.tibetan_size_pt >= options.meaning_size_pt + 6


@pytest.mark.parametrize("size", TIBETAN_SIZE_CHOICES)
def test_tibetan_size_choice_is_applied(size):
    document = build_document(make_rows(1), ExportOptions(tibetan_size_pt=size))
    run = tibetan_runs(document)[0]
    szCs = run._element.rPr.find(qn("w:szCs"))
    # w:szCs 的单位是半磅
    assert szCs.get(qn("w:val")) == str(int(size * 2))


def test_both_layouts_use_the_same_tibetan_size():
    """分栏排版不再把藏文收小——栏宽放得下。"""
    table = build_document(make_rows(1), ExportOptions(layout=LAYOUT_TABLE))
    columns = build_document(make_rows(1), ExportOptions(layout=LAYOUT_COLUMNS))

    def szcs(document):
        return tibetan_runs(document)[0]._element.rPr.find(qn("w:szCs")).get(qn("w:val"))

    assert szcs(table) == szcs(columns)


def test_rpr_children_follow_ooxml_order():
    """rPr 子元素顺序写错会导致 Word 报「文件已损坏」。"""
    document = build_document(make_rows(1), table_options())
    run = tibetan_runs(document)[0]
    tags = ["w:" + child.tag.split("}")[-1] for child in run._element.rPr]
    assert tags.index("w:rFonts") < tags.index("w:sz") < tags.index("w:szCs")


def test_exported_bytes_are_a_valid_docx():
    data = export_to_bytes(make_rows(3), table_options(title="测试"))
    document = Document(io.BytesIO(data))
    assert len(document.tables) == 1
    assert len(document.tables[0].rows) == 3


# ---------------------------------------------------------------------------
# 释义开关的三种组合
# ---------------------------------------------------------------------------
def test_include_meaning_gives_two_columns():
    document = build_document(
        make_rows(2, "太阳"),
        table_options(include_meaning=True, hide_meaning=False),
    )
    table = document.tables[0]
    assert len(table.columns) == 2
    assert table.cell(0, 1).text == "太阳0"


def test_hide_meaning_keeps_two_columns_but_blanks_them():
    document = build_document(
        make_rows(2, "太阳"),
        table_options(include_meaning=True, hide_meaning=True),
    )
    table = document.tables[0]
    assert len(table.columns) == 2
    assert all(row.cells[1].text.strip() == "" for row in table.rows)


def test_exclude_meaning_gives_one_column():
    document = build_document(
        make_rows(2, "太阳"),
        table_options(include_meaning=False, hide_meaning=False),
    )
    table = document.tables[0]
    assert len(table.columns) == 1
    assert table.cell(0, 0).text  # 藏文还在


def test_hide_meaning_without_include_meaning_still_one_column():
    document = build_document(
        make_rows(2), table_options(include_meaning=False, hide_meaning=True)
    )
    assert len(document.tables[0].columns) == 1


# ---------------------------------------------------------------------------
# 排版
# ---------------------------------------------------------------------------
def test_title_is_written_when_provided():
    with_title = build_document(make_rows(1), table_options(title="藏语基础词汇"))
    assert any(p.text == "藏语基础词汇" for p in with_title.paragraphs)

    without_title = build_document(make_rows(1), table_options(title=""))
    assert all(not p.text.strip() for p in without_title.paragraphs)


def test_no_header_row_by_design():
    """按需求，表格不加表头行，第一行直接就是第一个词条。"""
    document = build_document(make_rows(3, "太阳"), table_options())
    table = document.tables[0]
    assert table.cell(0, 0).text == "ཀ0"
    assert table.cell(0, 1).text == "太阳0"


def test_row_height_is_increased():
    document = build_document(make_rows(2), table_options(row_height_cm=1.2))
    for row in document.tables[0].rows:
        trHeight = row._tr.find(qn("w:trPr")).find(qn("w:trHeight"))
        assert trHeight is not None
        # 1.2cm = 680 twips
        assert int(trHeight.get(qn("w:val"))) == pytest.approx(680, abs=2)


def test_footer_contains_page_number_fields():
    document = build_document(make_rows(2), table_options())
    footer_xml = document.sections[0].footer.paragraphs[0]._p.xml
    assert "PAGE" in footer_xml
    assert "NUMPAGES" in footer_xml
    assert "第 " in footer_xml and " 页" in footer_xml


def test_empty_word_list_still_produces_a_document():
    data = export_to_bytes([], table_options(title="空表"))
    document = Document(io.BytesIO(data))
    assert document.tables == []
    assert any("没有符合条件的词条" in p.text for p in document.paragraphs)


# ---------------------------------------------------------------------------
# 分栏排版（默认）
# ---------------------------------------------------------------------------
def column_options(**overrides) -> ExportOptions:
    overrides.setdefault("layout", LAYOUT_COLUMNS)
    return ExportOptions(**overrides)


def test_columns_is_the_default_layout():
    options = ExportOptions()
    assert options.layout == LAYOUT_COLUMNS
    assert options.use_columns is True
    assert options.column_count == 3
    assert options.landscape is True


def test_columns_layout_has_no_table():
    document = build_document(make_rows(6), column_options())
    assert document.tables == [], "分栏排版不应该有表格"
    # 每个词条一个段落
    texts = [p.text for p in document.paragraphs if p.text.strip()]
    assert len(texts) == 6


def test_columns_layout_sets_w_cols_on_the_content_section():
    data = export_to_bytes(make_rows(6), column_options(column_count=3))
    xml = document_xml(data)

    cols = re.findall(r'<w:cols[^>]*w:num="(\d+)"[^>]*/>', xml)
    assert cols == ["3"], f"应当在正文那一节写入 w:cols，实际：{cols}"
    assert 'w:equalWidth="1"' in xml


def test_columns_layout_splits_title_and_body_into_two_sections():
    """标题必须独占整行宽度，所以要跟正文分属两节。

    否则 Word 会把标题也排进第一栏里，三栏的页面顶上会只有一栏有标题。
    """
    document = build_document(make_rows(6), column_options(title="藏语词汇"))
    assert len(document.sections) == 2, "分栏导出应当是「标题节 + 正文节」两节"

    title_section, body_section = document.sections
    # 第一节单栏，第二节多栏
    assert title_section._sectPr.find(qn("w:cols")) is None or \
        title_section._sectPr.find(qn("w:cols")).get(qn("w:num")) in (None, "1")
    assert body_section._sectPr.find(qn("w:cols")).get(qn("w:num")) == "3"


@pytest.mark.parametrize("count", [2, 3, 4])
def test_column_count_is_respected(count):
    xml = document_xml(export_to_bytes(make_rows(4), column_options(column_count=count)))
    assert f'w:num="{count}"' in xml


def test_columns_layout_landscape_swaps_page_size():
    landscape = document_xml(export_to_bytes(make_rows(2), column_options(landscape=True)))
    portrait = document_xml(export_to_bytes(make_rows(2), column_options(landscape=False)))

    # A4 横向：宽 16838 twips、高 11906；纵向反过来
    assert '<w:pgSz w:w="16838" w:h="11906" w:orient="landscape"/>' in landscape
    assert '<w:pgSz w:w="11906" w:h="16838"/>' in portrait


def test_columns_layout_meaning_visibility():
    visible = build_document(make_rows(2, "太阳"), column_options())
    texts = [p.text for p in visible.paragraphs if p.text.strip()]
    assert texts[0].startswith("ཀ0")
    assert "太阳0" in texts[0]

    hidden = build_document(
        make_rows(2, "太阳"), column_options(include_meaning=True, hide_meaning=True)
    )
    hidden_texts = [p.text for p in hidden.paragraphs if p.text.strip()]
    assert hidden_texts[0] == "ཀ0"          # 只剩藏文，自测用

    without = build_document(make_rows(2, "太阳"), column_options(include_meaning=False))
    assert all("太阳" not in p.text for p in without.paragraphs)


def test_columns_layout_uses_tibetan_font():
    runs = [run for p in build_document(make_rows(3), column_options()).paragraphs
            for run in p.runs if any(0x0F00 <= ord(c) <= 0x0FFF for c in run.text)]
    assert runs
    for run in runs:
        rFonts = run._element.rPr.find(qn("w:rFonts"))
        assert rFonts.get(qn("w:cs")) == DEFAULT_TIBETAN_FONT
        assert run._element.rPr.find(qn("w:szCs")) is not None


def test_columns_layout_empty_word_list_has_no_column_section():
    """没有词条时不该开分栏节——那句提示会被挤进一个很窄的栏里。"""
    document = build_document([], column_options(title="空"))
    assert len(document.sections) == 1
    assert any("没有符合条件的词条" in p.text for p in document.paragraphs)


def test_columns_layout_output_is_a_readable_docx():
    data = export_to_bytes(make_rows(20), column_options(title="词汇表"))
    document = Document(io.BytesIO(data))
    assert document.tables == []
    assert len([p for p in document.paragraphs if p.text.strip()]) >= 20


def test_document_contains_tibetan_text_and_grid_style():
    data = export_to_bytes(make_rows(2), table_options())
    xml = document_xml(data)
    assert "ཀ0" in xml
    assert "Table Grid" in xml or "TableGrid" in xml


def test_export_does_not_mutate_input_rows():
    rows = make_rows(10)
    original = [w.id for w in rows]
    apply_order(rows, table_options(order=ORDER_RANDOM, seed="1"))
    assert [w.id for w in rows] == original


def test_control_characters_do_not_break_the_export():
    """数据里混进控制字符时，导出不能整份失败。

    回归测试：控制字符（\\x0b、\\x00 之类）会先让列表页一切正常，
    一点「导出 Word」才抛 lxml 的 ValueError —— 用户完全不知道问题出在哪。
    写入时已经过滤过一次，但老数据库、手工改过的备份仍可能带进来，
    所以导出这条路必须自己兜住。
    """
    hostile = "x" + chr(11) + "y" + chr(0) + "z" + chr(31) + chr(127)
    rows = [models.Word(id=1, tibetan=hostile, meaning=hostile)]

    data = export_to_bytes(rows, table_options(title=hostile))
    assert data, "应当照常生成文档"

    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        for name in archive.namelist():
            if not name.endswith(".xml"):
                continue
            text = archive.read(name).decode("utf-8", errors="replace")
            for code in list(range(0, 9)) + [11, 12] + list(range(14, 32)) + [127]:
                assert chr(code) not in text, f"{name} 里残留控制字符 U+{code:04X}"

    # 正常文字被保留下来
    xml = document_xml(data)
    assert "xyz" in xml


def test_safe_text_helper():
    assert _safe_text("x" + chr(11) + "y") == "xy"
    assert _safe_text("བོད་ཡིག") == "བོད་ཡིག"
    assert _safe_text("") == ""
    assert _safe_text(None) == ""


def test_default_font_is_first_candidate():
    assert DEFAULT_TIBETAN_FONT == "Microsoft Himalaya"
    assert re.match(r"^[A-Za-z]", DEFAULT_TIBETAN_FONT)
